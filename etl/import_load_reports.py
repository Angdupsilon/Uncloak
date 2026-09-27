"""Build data/seed/dc_load_reports.csv: large-load and data-center load figures by region, each with
its scope, publisher, document and the quote or derivation behind it.

Scope (docs/us-expansion-plan.md rule 4: a data-center share is never assumed or carried across reports):
  data_centers             the publisher attributes the figure to data centers
  data_centers_and_crypto  the publisher's own category combines data centers and crypto mining
                           (Georgia Power from Q1 2026); never shown as data centers alone
  large_loads_all          all large loads; dc_share_pct is filled only when the same document
                           states the data-center share, with its quote

Sources:
  ERCOT  data/raw/ercot_large_load_queue.csv (unchanged rows; the two reports that state a
         data-center share, ~87% and ~90%, carry it)
  GA_PSC Georgia Power's quarterly Large Load Economic Development Reports (Georgia PSC Docket 55378/
         56002), data/raw/ga_psc_large_load/: per quarter, the announced load of the projects Georgia
         Power's attachment files under the data-center segment, all stages ("requested"), commitments
         = Contract for Electric Service + Request for Service ("committed", Georgia Power's own
         definition) and Contracts for Electric Service alone ("contracted")
  PJM    2026 PJM Load Forecast Report (posted 2026-01-14), data/raw/pjm_2026_load_forecast/:
         Table B-9 adjustments for the 12 zones PJM adjusted only for "Growth in data center load";
         for PS and DOM (adjusted for data centers plus another program) only the utilities' own
         data-center tables in their adjustment documents
Nothing is summed across publishers or regions.

Usage: python etl/import_load_reports.py
"""
from __future__ import annotations

import csv
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "seed" / "dc_load_reports.csv"
HEADER = ["ts", "region_key", "region_kind", "region_name", "metric", "scope", "value_mw", "forecast_year", "stage",
          "dc_share_pct", "dc_share_quote", "source_key", "source_url", "document", "quote"]

# ---------------------------------------------------------------- ERCOT
DC_SHARE = re.compile(r"of which ~(\d+)% are data centers")


def ercot_rows() -> list[dict]:
    rows = []
    with (RAW / "ercot_large_load_queue.csv").open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            share = DC_SHARE.search(r["quote_or_derivation"])
            for col, metric in (("gw_requested", "requested"), ("gw_approved", "approved"), ("gw_observed_peak", "observed_peak")):
                if not r[col]:
                    continue
                row = {"ts": r["ts"], "region_key": "ERCO", "region_kind": "ba",
                       "region_name": "Electric Reliability Council of Texas, Inc.", "metric": metric,
                       "scope": "large_loads_all", "value_mw": round(float(r[col]) * 1000, 6), "source_key": "ERCOT",
                       "source_url": r["source_url"], "document": r["document"], "quote": r["quote_or_derivation"]}
                if share and metric == "requested":
                    row["dc_share_pct"] = int(share.group(1))
                    row["dc_share_quote"] = share.group(0)
                rows.append(row)
    return rows


# ---------------------------------------------------------------- Georgia Power
GA_REGION = {"region_key": "GPC", "region_kind": "utility", "region_name": "Georgia Power Company"}
COMMITTED = {"contract for electric service", "request for service", "request for electric service"}


def stage_name(raw: str) -> str:
    """'Contract For Electric Service1' / '...**' -> 'contract for electric service' (footnote marks dropped)."""
    return re.sub(r"[\d*]+$", "", raw.strip()).strip().lower()


def ga_main(path: Path) -> list[dict]:
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Main"]
    rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    norm = lambda c: str(c).strip().lower() if c is not None else ""
    hi = next(i for i, r in enumerate(rows) if {"segment", "project stage"} <= {norm(c) for c in r})
    header = [norm(c) for c in rows[hi]]
    col: dict[str, int] = {}
    for i, k in enumerate(header):
        col.setdefault(k, i)  # later columns reuse names ("Project Stage" in the change section)
    load = next(i for i, k in enumerate(header) if k.startswith("announced load"))
    out = []
    for r in rows[hi + 1:]:
        stage, segment = r[col["project stage"]], r[col["segment"]]
        if stage and segment:
            out.append({"segment": str(segment).strip(), "stage": stage_name(str(stage)), "load": float(r[load] or 0)})
    return out


def ga_rows() -> list[dict]:
    rows = []
    with (RAW / "ga_psc_large_load" / "filings.csv").open(newline="", encoding="utf-8") as f:
        filings = list(csv.DictReader(f))
    for fl in filings:
        projects = ga_main(RAW / "ga_psc_large_load" / fl["attachment_file"])
        labels = {p["segment"] for p in projects if "data center" in p["segment"].lower()}
        if len(labels) != 1:
            raise SystemExit(f"{fl['quarter']}: expected one data-center segment label, got {labels}")
        label = labels.pop()
        scope = "data_centers_and_crypto" if "crypto" in label.lower() else "data_centers"
        dc = [p for p in projects if p["segment"] == label]
        unknown = {p["stage"] for p in dc} - COMMITTED - {"technical review"}
        if unknown:
            raise SystemExit(f"{fl['quarter']}: unknown project stages {unknown}")
        doc = f"Georgia Power Large Load Economic Development Report {fl['quarter'][:4]} Q{fl['quarter'][-1]} - Attachment (Main sheet)"
        base = {"ts": fl["period_end"], **GA_REGION, "scope": scope, "source_key": "GA_PSC", "source_url": fl["filing_url"],
                "document": doc + (f"; {fl['note']}" if fl["note"] else "")}
        groups = {
            "requested": (dc, "all stages"),
            "committed": ([p for p in dc if p["stage"] in COMMITTED], "Contract for Electric Service + Request for Service"),
            "contracted": ([p for p in dc if p["stage"] == "contract for electric service"], "Contract for Electric Service"),
        }
        for metric, (ps, stage) in groups.items():
            rows.append({**base, "metric": metric, "value_mw": round(sum(p["load"] for p in ps), 3), "stage": stage,
                         "quote": f"Sum of 'Announced Load*' over the {len(ps)} projects with Segment '{label}' "
                                  f"({stage}) as of {fl['period_end']}"})
    return rows


# ---------------------------------------------------------------- PJM
PJM_DIR = RAW / "pjm_2026_load_forecast"
PJM_REPORT = "https://www.pjm.com/-/media/DotCom/library/reports-notices/load-forecast/2026-load-report.pdf"
PJM_TABLES = "https://www.pjm.com/-/media/DotCom/planning/res-adeq/load-forecast/2026-load-report-tables.xlsx"
PJM_DATE = "2026-01-14"  # "2026 PJM Load Forecast Report, Posted Date: January 14, 2026"
DC_ONLY_ZONES = {
    "AEP": "American Electric Power", "ATSI": "American Transmission Systems, Inc.", "APS": "Allegheny Power Systems",
    "BGE": "Baltimore Gas and Electric", "COMED": "Commonwealth Edison", "DAYTON": "Dayton Power and Light",
    "DLCO": "Duquesne Light", "JCPL": "Jersey Central Power and Light", "METED": "Metropolitan Edison",
    "PECO": "PECO Energy", "PEPCO": "Potomac Electric Power", "PL": "PPL Electric Utilities",
}
PJM_ADJ_QUOTE = ("Load Report p.5, Load Adjustments: 'Zones Adjusted to account for: AEP, ATSI, APS, BGE, COMED, DAYTON, "
                 "DLCO, JCPL, METED, PECO, PEPCO, PL - Growth in data center load'; value from Table B-9 (updated 2/6/2026)")
# PSE&G adjustment document, p.5, "Table 1: PSE&G Data Center Peak Demand (MW) by Summer Year", Total column
# (values consider the completion rate).
PSEG_URL = "https://www.pjm.com/-/media/DotCom/planning/res-adeq/load-forecast/pseg-documentation.pdf"
# Dominion Energy letter to PJM (January 6, 2026), p.1: the requested data-center coincident peak by year. The
# table is an image in the PDF; values were read from the rendered page. The 2046 value matches the letter's text,
# "The Company is forecasting 16.6 GW of demand by 2046".
DOM_URL = "https://www.pjm.com/-/media/DotCom/planning/res-adeq/load-forecast/dominion-documentation.pdf"
DOM_FORECAST = {2026: 4433, 2027: 4933, 2028: 5448, 2029: 5980, 2030: 6526, 2031: 7086, 2032: 7661, 2033: 8247,
                2034: 8846, 2035: 9456, 2036: 10076, 2037: 10705, 2038: 11342, 2039: 11987, 2040: 12639, 2041: 13297,
                2042: 13959, 2043: 14625, 2044: 15294, 2045: 15965, 2046: 16636}


def pjm_rows() -> list[dict]:
    import pdfplumber

    rows = []
    ws = openpyxl.load_workbook(PJM_DIR / "2026-load-report-tables.xlsx", read_only=True, data_only=True)["Table B9"]
    table = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    years = next(r for r in table if r[1] == 2026)
    for r in table:
        zone = str(r[0]).strip() if r[0] else ""
        if zone in DC_ONLY_ZONES:
            for i, year in enumerate(years):
                if isinstance(year, int) and r[i] is not None:
                    rows.append({"ts": PJM_DATE, "region_key": f"PJM-{zone}", "region_kind": "zone",
                                 "region_name": f"PJM {zone} zone ({DC_ONLY_ZONES[zone]})", "metric": "forecast_adjustment",
                                 "scope": "data_centers", "value_mw": round(float(r[i]), 3), "forecast_year": year,
                                 "stage": "summer peak adjustment above embedded", "source_key": "PJM", "source_url": PJM_TABLES,
                                 "document": "2026 PJM Load Forecast Report, Table B-9", "quote": PJM_ADJ_QUOTE})
    with pdfplumber.open(PJM_DIR / "pseg-documentation.pdf") as pdf:
        text = pdf.pages[4].extract_text()
    for m in re.finditer(r"^(20[2-4]\d) 394 (\d+) (\d+) (\d+) (\d+)$", text, re.M):
        rows.append({"ts": PJM_DATE, "region_key": "PJM-PS", "region_kind": "zone",
                     "region_name": "PJM PS zone (Public Service Electric and Gas)", "metric": "forecast", "scope": "data_centers",
                     "value_mw": int(m.group(5)), "forecast_year": int(m.group(1)), "stage": "data-center summer peak demand",
                     "source_key": "PJM", "source_url": PSEG_URL, "document": "PSE&G 2026 Load Forecast Adjustments (January 2026), p.5",
                     "quote": "Table 1: PSE&G Data Center Peak Demand (MW) by Summer Year, Total (values consider the completion rate)"})
    dom = {"region_key": "PJM-DOM", "region_kind": "zone", "region_name": "PJM DOM zone (Dominion Virginia Power)",
           "scope": "data_centers", "source_key": "PJM", "source_url": DOM_URL,
           "document": "Dominion Energy letter to PJM Load Analysis Team, January 6, 2026"}
    for year, mw in DOM_FORECAST.items():
        rows.append({**dom, "ts": PJM_DATE, "metric": "forecast", "value_mw": mw, "forecast_year": year,
                     "stage": "requested data-center coincident peak",
                     "quote": "p.1 table image (read from the rendered page): 'The Company requested the coincident peak values "
                              "shown in the table to the left for its data center industry'"})
    rows.append({**dom, "ts": "2025-12-31", "metric": "observed_peak", "value_mw": 4000, "stage": "2025 data-center coincident peak",
                 "quote": "'the Company currently serves the largest data center market in the world with a 2025 coincident peak of 4 GW'"})
    for stage, gw in (("ESA (connected)", 9.8), ("CLOA (under construction)", 7.1), ("ELOA (engineering study)", 30.1)):
        rows.append({**dom, "ts": "2025-07-31", "metric": "contracted", "value_mw": round(gw * 1000), "stage": stage,
                     "quote": "p.3: 'as of July 2025, the capacity value of these contracts is 47 GW (9.8 ESA + 7.1 CLOA + 30.1 ELOA)'"})
    return rows


def main() -> None:
    rows = ercot_rows() + ga_rows() + pjm_rows()
    rows.sort(key=lambda r: (r["source_key"], r["region_key"], r["metric"], str(r.get("forecast_year") or ""), r["ts"], r.get("stage") or ""))
    with OUT.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=HEADER, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    by = defaultdict(int)
    for r in rows:
        by[(r["source_key"], r["scope"])] += 1
    print(f"{len(rows)} load-report rows -> {OUT.relative_to(ROOT)}: " + ", ".join(f"{s}/{sc} {n}" for (s, sc), n in sorted(by.items())))


if __name__ == "__main__":
    main()
